from __future__ import annotations

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "参赛材料"
ASSETS = OUT / "assets"
SCENE = ROOT / "public" / "scene"
ASSETS.mkdir(parents=True, exist_ok=True)

GREEN = "173F35"
MID_GREEN = "2E6A57"
PALE = "EEF4EA"
LIME = "C7F06C"
INK = "17221E"
MUTED = "5D6C66"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def font(size: int, bold: bool = False):
    names = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    for path in names:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rounded(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_visual_assets():
    lake = Image.open(SCENE / "swiss-lake.jpg").convert("RGB")
    cast = Image.open(SCENE / "login-cast.png").convert("RGBA")
    success = Image.open(SCENE / "login-success.png").convert("RGBA")

    hero = ImageOps.fit(lake, (1600, 820), method=Image.Resampling.LANCZOS)
    overlay = Image.new("RGBA", hero.size, (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    for x in range(hero.width):
        alpha = int(170 * (1 - x / hero.width) + 15)
        od.line((x, 0, x, hero.height), fill=(10, 44, 35, alpha))
    hero = Image.alpha_composite(hero.convert("RGBA"), overlay)
    scale = min(920 / cast.width, 760 / cast.height)
    cast_resized = cast.resize((int(cast.width * scale), int(cast.height * scale)), Image.Resampling.LANCZOS)
    hero.alpha_composite(cast_resized, (1600 - cast_resized.width + 40, 820 - cast_resized.height + 35))
    hero.save(ASSETS / "cover-hero.png")

    lineup = Image.new("RGB", (1600, 520), "#EEF4EA")
    d = ImageDraw.Draw(lineup)
    d.ellipse((-140, -210, 560, 490), fill="#DDEBCE")
    d.ellipse((1210, 80, 1780, 650), fill="#D9E8E1")
    scale = min(1120 / success.width, 470 / success.height)
    group = success.resize((int(success.width * scale), int(success.height * scale)), Image.Resampling.LANCZOS)
    lineup.paste(group, (820 - group.width // 2, 520 - group.height), group)
    d.text((58, 48), "一套角色贯穿整个体验", font=font(44, True), fill="#173F35")
    d.text((60, 112), "登录选择  画像对话  候选席位  圆桌发言  桌后关系", font=font(24), fill="#547067")
    rounded(d, (58, 392, 516, 466), 26, "#173F35")
    d.text((94, 411), "选择你喜欢的样子再开始说话", font=font(22, True), fill="#F6F8ED")
    lineup.save(ASSETS / "character-system.png")

    flow = Image.new("RGB", (1600, 500), "white")
    d = ImageDraw.Draw(flow)
    d.text((56, 36), "最小闭环", font=font(34, True), fill="#17221E")
    d.text((240, 46), "AI 只在关键节点工作  真人提供问题  经历与关系", font=font(20), fill="#5D6C66")
    nodes = [
        ("01", "选择身份", "一键进入"),
        ("02", "对话认识", "生成可编辑标签"),
        ("03", "互补匹配", "不是复制同类"),
        ("04", "真人圆桌", "Agent 主持不代答"),
        ("05", "收桌延续", "总结与新桌友"),
    ]
    xs = [55, 355, 655, 955, 1255]
    for idx, ((num, title, note), x) in enumerate(zip(nodes, xs)):
        rounded(d, (x, 125, x + 245, 390), 28, "#F0F5ED", "#D4DED6", 2)
        rounded(d, (x + 22, 148, x + 80, 206), 20, "#173F35")
        d.text((x + 36, 160), num, font=font(18, True), fill="#C7F06C")
        d.text((x + 22, 244), title, font=font(30, True), fill="#173F35")
        lines = note.split(" ") if " " in note else [note]
        d.text((x + 22, 306), "\n".join(lines), font=font(19), fill="#61736C", spacing=8)
        if idx < len(nodes) - 1:
            d.line((x + 250, 258, x + 292, 258), fill="#6A8E7D", width=4)
            d.polygon([(x + 292, 258), (x + 279, 249), (x + 279, 267)], fill="#6A8E7D")
    flow.save(ASSETS / "product-flow.png")

    arch = Image.new("RGB", (1600, 720), "white")
    d = ImageDraw.Draw(arch)
    d.text((60, 36), "Agent 编排与知乎能力接入位置", font=font(34, True), fill="#17221E")
    current = [
        (75, 170, 330, 300, "用户", "真实问题与经历"),
        (395, 145, 720, 325, "画像 Agent", "自由对话  标签确认"),
        (785, 145, 1110, 325, "匹配编排器", "话题交集  视角互补"),
        (1175, 145, 1515, 325, "四人圆桌", "真人表达  Agent 主持"),
        (785, 420, 1110, 610, "收桌 Agent", "共识  分歧  新问题"),
        (1175, 420, 1515, 610, "关系延续", "桌友  消息  再次相遇"),
    ]
    for x1, y1, x2, y2, title, note in current:
        rounded(d, (x1, y1, x2, y2), 24, "#EEF4EA", "#CCDAD0", 2)
        d.text((x1 + 24, y1 + 30), title, font=font(28, True), fill="#173F35")
        d.text((x1 + 24, y1 + 84), note, font=font(19), fill="#5D6C66")
    arrows = [((330,235),(395,235)),((720,235),(785,235)),((1110,235),(1175,235)),((1345,325),(950,420)),((1110,515),(1175,515))]
    for start, end in arrows:
        d.line((*start, *end), fill="#507A69", width=4)
    rounded(d, (75, 430, 720, 632), 24, "#173F35")
    d.text((105, 458), "知乎开放能力  待确认后接入", font=font(25, True), fill="#F4F7E9")
    d.text((105, 515), "OAuth  搜索  热榜  关注流  直答 Agent", font=font(20), fill="#D9E9C1")
    d.text((105, 558), "只提供身份  话题与依据  不替用户表达立场", font=font(18), fill="#AFC4B6")
    d.line((720, 530, 785, 530), fill="#C7F06C", width=4)
    arch.save(ASSETS / "architecture.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), LINE)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_font(run, size=11, bold=False, color=INK, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, 8, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, 8, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, 18 if level == 1 else 13, True, "000000", "Microsoft YaHei")
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.45
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 11, True)
        text = text[len(bold_lead):]
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_font(p.add_run(item), 10.2)


def add_table(doc, headers, rows, widths=None, small=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, GREEN)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), 9.2 if small else 9.8, True, WHITE)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[idx])
            if row_idx % 2:
                set_cell_shading(cell, "F5F8F3")
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            if idx == 0 and len(headers) > 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(value)), 8.7 if small else 9.3, bold=(idx == 0), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    set_font(cap.add_run(caption), 8.5, color=MUTED)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_doc():
    make_visual_assets()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(.62)
    sec.bottom_margin = Inches(.62)
    sec.left_margin = Inches(.7)
    sec.right_margin = Inches(.7)
    sec.header_distance = Inches(.25)
    sec.footer_distance = Inches(.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for name, size in (("Title", 31), ("Heading 1", 18), ("Heading 2", 13)):
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
    if "Figure Caption" not in [s.name for s in styles]:
        styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)

    header = sec.header.paragraphs[0]
    header.text = "组一桌  知乎黑客松 2026  灵魂匹配局"
    set_font(header.runs[0], 8, True, MUTED)
    add_page_number(sec.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("组一桌 知乎黑客松参赛计划书"), 31, True, "000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run("灵魂匹配局  社区连接与兴趣社交"), 13, True, MID_GREEN)
    add_figure(doc, ASSETS / "cover-hero.png", 7.08, "瑞士湖边视觉与四位原创角色  登录选择后进入同一套 3D 社交空间")
    add_body(doc, "组一桌用 Agent 降低高质量真人讨论的组局成本。AI 不替用户回答，而是在认识用户、补齐视角、主持分歧和沉淀关系四个节点工作，让一次真实问题从“有人回应”走向“形成新认识”，再走向“知道为什么还想再见”。")
    add_table(doc, ["当前赛道", "提交定位", "本版状态"], [["赛道一", "社区连接与兴趣社交", "最小 Demo 闭环已跑通"]], [1.1, 3.1, 2.2])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("版本 2026 年 9 月 15 日"), 8.5, color=MUTED)

    page_break(doc)
    add_heading(doc, "一 项目主张", 1)
    add_body(doc, "结论先行  组一桌不是新奇的 3D 聊天室，也不是四个 Agent 自说自话。它解决的是一个常见但长期没有被处理好的问题：当用户真正想听见不同人的经验时，很难在恰当时间找到互补的人、开启不尴尬的讨论，并把这次相遇延续下去。", "结论先行  ")
    add_heading(doc, "真实痛点", 2)
    add_bullets(doc, [
        "内容平台擅长提供答案，却不一定能让提问者遇见愿意继续聊的人。",
        "传统兴趣匹配偏向相似度，容易复制同温层；好讨论需要事实、经验、理论、追问等互补贡献。",
        "陌生人开场成本高，讨论也容易散掉；结束后没有问题进化、依据和关系记忆。",
        "评审现场很难同时找到四位真人，若 Demo 依赖在线人数，完整闭环不可重复。",
    ])
    add_heading(doc, "产品价值链", 2)
    add_figure(doc, ASSETS / "product-flow.png", 7.05, "从身份选择到关系延续的完整闭环  AI 工作于关键节点  不替真人形成立场")
    add_body(doc, "产品价值不止是“组到一桌”。它同时减少三类成本：找到合适人的搜索成本、让陌生人进入有效讨论的协调成本、让知识和关系留下来的沉淀成本。")

    page_break(doc)
    add_heading(doc, "二 评委可完整体验的用户流程", 1)
    add_figure(doc, ASSETS / "character-system.png", 7.0, "原创 IP 形象贯穿登录  对话  候选人  圆桌与收桌  降低陌生社交的心理压力")
    add_table(doc, ["阶段", "用户看到什么", "评审证据"], [
        ("一键进入", "湖边场景与四个 3D 人物  选择一个角色即进入", "设计完成度与角色一致性"),
        ("Agent 认识我", "两轮自由聊天  实时追问而非固定问卷", "阶跃星辰真实生成"),
        ("标签确认", "身份  经历  兴趣  表达视角可编辑和隐藏", "匹配可解释  隐私可控"),
        ("开车找桌", "SUV 沿地图驶向最佳匹配桌  可跳过", "3D 不是装饰  映射产品逻辑"),
        ("四人圆桌", "用户真实发言  三位互补角色回应  阿桌主持", "实时性  互补性  主持边界"),
        ("收桌卡片", "共识  分歧  用户贡献  新问题  值得继续认识的人", "从讨论到关系的闭环"),
    ], [1.0, 3.0, 2.6], small=True)
    add_body(doc, "冷启动处理  评审 Demo 使用三位明确标注的虚构桌友，保证无人在线时也能完整体验；上桌后的新回复由模型针对评委真实发言实时生成。正式产品用真人席位替换预置角色，Agent 仍只负责匹配、主持和整理。", "冷启动处理  ")

    page_break(doc)
    add_heading(doc, "三 演示桌设计", 1)
    add_body(doc, "首桌主题是“离开大城市，是逃避还是重新选择生活”。它比讨论 AI 是否替代专业判断更能体现真人社交：任何用户都能带着生活经验加入，三位桌友的差异来自处境和取舍，而不是模型在讨论模型。")
    add_table(doc, ["角色", "身份牌", "带来的贡献", "不做什么"], [
        ("用户", "由前置聊天生成", "真实处境  犹豫与问题", "不要求先有成熟观点"),
        ("林夏", "县城医生  亲历视角", "离开后的获得与代价", "不把返乡包装成田园童话"),
        ("周砚", "产品工程师  机会视角", "关系网络  机会成本  低成本试验", "不把留下等同于妥协"),
        ("程野", "城市研究者  结构视角", "住房  照护  公共服务与掌控感", "不把选择道德化"),
        ("阿桌", "圆桌主持 Agent", "递话  追问  分歧重构  收桌", "不代答  不投票  不伪造共识"),
    ], [0.8, 1.6, 2.65, 1.65], small=True)
    add_heading(doc, "讨论如何产生新价值", 2)
    add_body(doc, "原问题把选择压缩成“逃避或勇敢”。圆桌通过不同人的真实约束，把它推进为一个可行动的新问题：“你真正想离开的，是城市，还是一种失去掌控的生活？”用户随后可以继续验证关系能否迁移、机会能否保留、压力源是否会随迁居消失。")
    add_body(doc, "社交结果不是简单互关。收桌卡解释用户为何还想认识某个人，并把选中的桌友带回好友栏，留下下一次对话的上下文。")

    page_break(doc)
    add_heading(doc, "四 AI 系统与产品边界", 1)
    add_figure(doc, ASSETS / "architecture.png", 7.05, "当前闭环与知乎开放能力的推荐接入位置  深绿色部分为待确认方案")
    add_table(doc, ["Agent 能力", "输入", "输出", "价值与约束"], [
        ("画像 Agent", "两轮自由表达", "结构化标签与身份牌", "用户确认后才保存  不读取知乎私密数据"),
        ("匹配编排", "话题  经历  表达视角", "三位互补候选人和理由", "避免只按相似度扩大信息茧房"),
        ("桌友生成", "角色边界  桌上原话", "针对用户发言的新回复", "Demo 角色明确标注虚构  不冒充真人"),
        ("主持 Agent", "当前分歧与节奏", "递话  追问  重构  收束", "优先沉默  不抢表达权"),
        ("收桌 Agent", "实际消息与证据轮次", "共识  分歧  贡献  新问题", "原话优先  不把个人意见写成全桌共识"),
    ], [1.05, 1.65, 1.65, 2.35], small=True)
    add_body(doc, "当前技术实现包括 React 前端、Three.js 和 GLB 场景、FastAPI 后端、WebSocket 桌内事件、阶跃星辰兼容接口、本地持久化与可恢复的收桌产物。API 密钥只在后端本地环境中读取，不进入浏览器和代码仓库。")

    page_break(doc)
    add_heading(doc, "五 对照官方规则的踩分策略", 1)
    add_body(doc, "主办方初审单件作品看材料、Demo 与打分的总时间控制在 3 分钟以内；同分时先比较 AI 场景价值，再比较完成度。因此首页第一屏、计划书前两页和演示前 30 秒必须先证明场景成立，再展示技术。")
    add_table(doc, ["官方维度", "组一桌的直接证据", "演示方式"], [
        ("AI 场景价值 40%", "降低真人高质量讨论的搜索  协调与沉淀成本  适配知乎真实问题与专业讨论", "首页一句话  前置画像  收桌关系闭环"),
        ("创新度 25%", "不按相似度复制同类  用贡献互补组桌  Agent 以克制主持而非代答", "匹配理由  身份牌  阿桌主持边界"),
        ("完成度", "从登录到好友回流可重复跑通  等待  失败  重试和边界有反馈", "全程实操  不切 PPT 假画面"),
        ("产品体验与设计 10%", "湖边低压空间  原创 IP  统一人物形象  3D 车与场景服务于路径", "横屏录制  字体与 GLB 预热"),
    ], [1.25, 3.3, 2.15], small=True)
    add_heading(doc, "三分钟内必须讲清的四件事", 2)
    add_bullets(doc, [
        "我们连接真人，不让 Agent 替真人聊天；模拟桌友只解决评审冷启动，并清楚标注。",
        "画像来自自由对话，标签可改；匹配依据是话题交集、视角互补和可分享的经历。",
        "用户说一句，桌友针对这句话实时回应；阿桌只在需要时介入。",
        "讨论结束会留下更好的问题和继续认识的人，产品不止于“组一桌”。",
    ])

    page_break(doc)
    add_heading(doc, "六 知乎资源接入方案", 1)
    add_body(doc, "本节是待确认方案，当前代码尚未接入。优先选择能直接证明“生长在知乎生态里”的能力，不为了堆接口破坏产品逻辑。官方 hackathon skill 是开发工具，用来保证鉴权、参数和限额使用正确；仅安装 skill 本身不会形成用户价值，必须把 API 结果变成可见体验。")
    add_table(doc, ["优先级", "知乎能力", "在组一桌中的用法", "为什么值得做"], [
        ("P0", "知乎 OAuth", "替换正式版一键身份入口  头像昵称只在授权后读取", "官方推荐  登录人数也是人气奖信号之一"),
        ("P0", "知乎搜索", "根据桌主题检索高相关  高权威内容  主持在事实分歧时投放可追溯依据卡", "最贴合知乎知识资产  直接提升讨论质量"),
        ("P1", "知乎热榜", "每日生成可组的热门桌  服务端缓存  避免重复调用", "让桌子跟随社区正在发生的问题"),
        ("P1", "关注流与关系", "在用户同意后发现共同关注和弱连接  作为匹配解释的一部分", "把一次匹配接到真实社区关系"),
        ("P1", "直答 Agent", "仅在主持需要澄清事实时生成有依据的短答  不参与立场表达", "为分歧补事实  不把产品变成问答机器人"),
        ("P2", "知乎知识", "把主题背景整理成桌前 30 秒知识卡  供所有人对齐语境", "降低专业话题的入桌门槛"),
        ("可选", "刘看山 IP", "只做入口彩蛋或官方能力提示  不替换阿桌原创形象", "保留比赛氛围  同时避免稀释产品品牌"),
    ], [.65, 1.2, 3.35, 1.5], small=True)
    add_body(doc, "推荐最小接入组合  OAuth 加知乎搜索。它同时覆盖人气、真实账号入口和知乎内容价值。若时间只够一个内容接口，先做搜索依据卡；热榜与关注流放到决赛迭代。所有官方数据在服务端缓存，并在 UI 显示来源；禁止批量爬取和未经授权读取用户隐私。", "推荐最小接入组合  ")

    page_break(doc)
    add_heading(doc, "七 演示视频与提交流程", 1)
    add_table(doc, ["时间", "画面", "旁白重点"], [
        ("0 至 12 秒", "湖边登录选择角色  进入首页", "我们不是替用户回答  而是让不同的人一起判断"),
        ("12 至 45 秒", "和阿桌自由聊两轮  生成并确认标签", "不是固定问卷  用户决定哪些标签公开"),
        ("45 至 62 秒", "互补候选人与 SUV 找桌", "不是随机拼桌  也不是复制一个相似的你"),
        ("62 至 108 秒", "用户发言  桌友实时回应", "三位预置角色只解决冷启动  新回复由模型实时生成"),
        ("108 至 135 秒", "生成收桌卡  选择继续认识", "AI 让一次讨论变成可延续的人际连接"),
    ], [.9, 2.85, 2.95], small=True)
    add_heading(doc, "提交前清单", 2)
    add_table(doc, ["材料", "规则", "当前状态与下一步"], [
        ("公网 Demo", "必交  评委可直接完整体验", "待部署  配置后端密钥与健康检查  用新设备无痕验收"),
        ("产品计划书", "必交  初审重点", "本文件已完成  提交前补团队名与公网地址"),
        ("代码仓库", "选交加分", "已有 GitHub 仓库  提交前清理密钥  补 README 与启动说明"),
        ("演示视频", "选交加分", "按上表录制  控制在 2 分 15 秒左右  加字幕和边界说明"),
        ("知乎想法", "人气入口", "发布项目问题和 20 秒动图  绑定官方话题与圈子  禁止刷量"),
    ], [1.1, 2.0, 3.6], small=True)
    add_body(doc, "最优时间顺序  先保证公网链接从陌生设备能完整跑通，再录制视频；随后接 OAuth 与知乎搜索的最小闭环。不要在提交前同时重做多人数据库、复杂 3D 动画和多个 API。")

    page_break(doc)
    add_heading(doc, "八 当前完成度与风险", 1)
    add_table(doc, ["模块", "当前结果", "提交前判定"], [
        ("一键角色登录", "已完成  四角色可选  选择后候选人顺延去重", "可录制"),
        ("自由对话画像", "已完成  阶跃星辰实时追问  标签可编辑和隐藏", "可录制"),
        ("互补匹配", "已完成演示解释  正式真人数据仍待接入", "Demo 可用  需披露"),
        ("3D 路线与圆桌", "已完成  湖边视觉  GLB 人物  SUV 与圆桌空间保留", "可录制"),
        ("实时圆桌", "已完成  用户消息触发角色新回应  WebSocket 状态可见", "可录制"),
        ("总结与社交结果", "已完成并实机验证  等待状态  卡片  桌友回流", "可录制"),
        ("公网与正式账号", "尚未完成", "阻断提交  最高优先级"),
        ("知乎官方 API", "方案已定  尚未接入", "OAuth 与搜索优先  需用户确认"),
    ], [1.25, 3.2, 2.25], small=True)
    add_heading(doc, "主要风险与应对", 2)
    add_bullets(doc, [
        "生成时间波动  收桌立即显示阿桌进度，并在 10 秒、24 秒和 45 秒主动恢复产物；失败后允许重试。",
        "模拟角色被误解为真人  入口、人物资料和提交材料统一标注虚构经历与实时生成边界。",
        "3D 资源较重  公网部署前开启压缩与缓存，首屏先展示湖景与角色图，GLB 延后加载。",
        "知乎 API 限额  热榜和直答按官方限制做服务端缓存；搜索按主题去重，不在每次渲染时请求。",
        "真实多人尚未完成  本次初审只承诺可重复的体验闭环；决赛路线再扩展实时真人房间与关系数据库。",
    ])

    page_break(doc)
    add_heading(doc, "九 评委可能追问", 1)
    add_table(doc, ["问题", "建议回答"], [
        ("为什么不用 AI 直接回答", "答案解决信息需求  组一桌解决的是人与人的理解  分歧和关系延续  Agent 只降低组局成本"),
        ("现在不还是 Agent 模拟人吗", "评审冷启动用三位明确标注的虚构角色保证可重复  正式产品四个席位是真人  模拟回复只证明调度和主持闭环"),
        ("匹配和推荐有什么不同", "推荐把内容给一个人  组一桌要同时优化四个人的贡献组合  目标是互补而非相似"),
        ("如何避免信息茧房", "匹配显式保留事实  经验  理论和追问等不同贡献  用户可看到并修改匹配依据"),
        ("知乎为什么需要它", "知乎已有真实问题  专业内容和关系网络  组一桌把内容消费推进为小规模讨论  再把新问题和关系带回社区"),
        ("隐私如何处理", "只保存用户确认的标签  OAuth 后最小授权  私密数据不进入桌面  来源卡只引用公开内容"),
    ], [2.0, 4.7], small=True)
    add_heading(doc, "最终定位", 2)
    add_body(doc, "组一桌的竞争力不是“好玩”或“新颖”本身，而是一个可验证的社区价值闭环：把真实问题带到合适的人之间，让差异变得可讨论，让讨论留下知识结构和关系理由。3D 场景和原创角色降低表达压力，Agent 在关键节点承担组织成本，知乎内容与关系能力则让产品从 Demo 生长为社区基础设施。")
    add_body(doc, "仓库  https://github.com/JiuTian-dev/Zu-Yi-Zhuo")
    add_body(doc, "规则依据  知乎黑客松 2026 校园新锐季开发者手册  FAQ  参赛者开发流程文档  核对日期 2026 年 9 月 15 日")

    out = OUT / "组一桌_知乎黑客松参赛计划书.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build_doc()
